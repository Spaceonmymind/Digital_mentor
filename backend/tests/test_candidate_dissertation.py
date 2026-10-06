import io
import asyncio
import json

import fitz
import pytest
from docx import Document as DocxDocument
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from app.db.models import Analysis, AnalysisResult, Document
from app.db.session import async_session_factory
from app.execution.candidate_dissertation import CandidateDissertationAgentFlow
from app.execution.rule_schemas import CandidateAgentOutput, RuleCheckResult, RuleEvidence
from app.execution.rule_scoring import score_criterion, score_methodology
from app.methodology.models import Methodology, MethodologyCriterion
from app.assessment.models import Assessment
from app.execution.rule_schemas import CandidateFinalOutput
from app.llm.schemas import LLMResult, LLMUsage
from app.methodology.registry import MethodologyRegistry
from app.methodology.repository import MethodologyRepository
from app.methodology.seeds.candidate_dissertation import ensure_candidate_dissertation_seed
from app.methodology.seeds.candidate_dissertation.data_v1 import DISCLAIMER
from app.methodology.seeds.startup_vkr import ensure_startup_vkr_seed
from app.services.extraction import TextExtractionService
from app.services.reports import ReportService


def rule(status, *, strength="REQUIRED", weight=1.0):
    return RuleCheckResult(
        rule_code="CD1.R01", criterion_code="CD1", title="Правило", status=status,
        finding="Результат", recommendation=None, confidence=1, evidence=[],
        source_document="Методические рекомендации", source_section="2.1", source_pages=[13],
        capability="TEXT", normative_strength=strength, score_weight=weight,
    )


@pytest.mark.asyncio
async def test_candidate_methodology_is_available_and_isolated_from_startup(client):
    async with async_session_factory() as session:
        candidate = await ensure_candidate_dissertation_seed(session)
        startup = await ensure_startup_vkr_seed(session)
        loaded = (
            await session.execute(
                select(Methodology).where(Methodology.id == candidate.id).options(
                    selectinload(Methodology.criteria).selectinload(MethodologyCriterion.indicators),
                    selectinload(Methodology.agents),
                )
            )
        ).scalar_one()
        startup_loaded = (
            await session.execute(
                select(Methodology).where(Methodology.id == startup.id).options(selectinload(Methodology.criteria))
            )
        ).scalar_one()
        candidate_codes = {item.number for item in loaded.criteria}
        startup_codes = {item.number for item in startup_loaded.criteria}

    response = await client.get("/api/v1/methodologies")
    items = response.json()
    candidate_item = next(item for item in items if item["work_type"] == "CANDIDATE_DISSERTATION")
    assert candidate_item["availability"] == "AVAILABLE"
    assert candidate_item["active_version"] == "1.0"
    assert candidate_item["max_score"] == 70
    assert candidate_codes == {f"CD{i}" for i in range(1, 8)}
    assert startup_codes == {f"C{i}" for i in range(1, 7)}
    assert candidate.configuration["execution_profile"] == "candidate_dissertation"
    assert startup.configuration["execution_profile"] == "startup_vkr"
    thematic_agents = [item for item in loaded.agents if item.code != "A-01"]
    assert {item.code for item in thematic_agents} == {"CD-10", "CD-20", "CD-30", "CD-40", "CD-50", "CD-60"}
    assert all(item.execution_mode == "parallel" for item in thematic_agents)
    assert next(item for item in loaded.agents if item.code == "A-01").execution_mode == "final"
    assert all(item["availability"] == "NOT_CONFIGURED" for item in items if item["work_type"] not in {"STARTUP_VKR", "CANDIDATE_DISSERTATION"})


@pytest.mark.asyncio
async def test_registry_resolves_candidate_and_startup_versions():
    async with async_session_factory() as session:
        await ensure_candidate_dissertation_seed(session)
        await ensure_startup_vkr_seed(session)
        registry = MethodologyRegistry(MethodologyRepository(session))
        candidate = await registry.resolve("CANDIDATE_DISSERTATION")
        startup = await registry.resolve("STARTUP_VKR")
    assert candidate.version == "1.0"
    assert startup.version == "2.0"
    assert candidate.max_score == 70
    assert startup.max_score == 60


@pytest.mark.asyncio
async def test_candidate_rules_keep_applications_only_in_cd1_and_trace_sources():
    async with async_session_factory() as session:
        candidate = await ensure_candidate_dissertation_seed(session)
        loaded = (
            await session.execute(
                select(Methodology).where(Methodology.id == candidate.id).options(
                    selectinload(Methodology.criteria).selectinload(MethodologyCriterion.indicators)
                )
            )
        ).scalar_one()
    by_code = {item.number: item for item in loaded.criteria}
    assert any("прилож" in rule.title.lower() for rule in by_code["CD1"].indicators)
    assert not any("прилож" in rule.title.lower() for rule in by_code["CD6"].indicators)
    assert [len(by_code[f"CD{i}"].indicators) for i in range(1, 8)] == [17, 13, 3, 4, 5, 10, 15]
    for criterion in loaded.criteria:
        for indicator in criterion.indicators:
            config = indicator.configuration
            assert config["source_document"]
            assert config["source_section"]
            assert config["source_pages"]
            assert config["capability"]
            assert config["normative_strength"] in {"REQUIRED", "CONDITIONAL", "RECOMMENDED", "OPTIONAL"}


def test_rule_scoring_excludes_not_checked_not_applicable_optional_and_recommended():
    calculated = score_criterion([
        rule("PASS"), rule("PARTIAL"), rule("FAIL"), rule("NOT_CHECKED"),
        rule("NOT_APPLICABLE"), rule("FAIL", strength="OPTIONAL"), rule("FAIL", strength="RECOMMENDED"),
    ])
    assert calculated.score == 5
    assert calculated.max_score == 10
    empty = score_criterion([rule("NOT_CHECKED"), rule("NOT_APPLICABLE")])
    assert empty.score is None
    assert empty.max_score is None
    assert empty.status == "NOT_CHECKED"


def test_methodology_scoring_has_dynamic_evaluated_max_and_separate_coverage():
    result = score_methodology({
        "CD1": [rule("PASS"), rule("NOT_CHECKED")],
        "CD2": [rule("NOT_CHECKED")],
    })
    assert result["overall_score"] == 10
    assert result["evaluated_max_score"] == 10
    assert result["nominal_max_score"] == 20
    assert result["coverage"] == pytest.approx(1 / 3, abs=0.0001)


@pytest.mark.asyncio
async def test_plain_text_pdf_formatting_rules_are_not_checked_not_failed():
    async with async_session_factory() as session:
        candidate = await ensure_candidate_dissertation_seed(session)
        loaded = (
            await session.execute(
                select(Methodology).where(Methodology.id == candidate.id).options(
                    selectinload(Methodology.criteria).selectinload(MethodologyCriterion.indicators)
                )
            )
        ).scalar_one()
        cd7 = next(item for item in loaded.criteria if item.number == "CD7")
        checked, pending = CandidateDissertationAgentFlow(session)._precheck_rules(cd7.indicators, {"full_text": "Текст"}, "pdf")
    by_code = {item.rule_code: item for item in checked}
    assert by_code["CD7.R03"].status == "NOT_CHECKED"
    assert by_code["CD7.R11"].status == "NOT_CHECKED"
    assert by_code["CD7.R13"].status == "NOT_CHECKED"
    assert all(item.status != "FAIL" for item in checked)
    assert any(item.configuration["capability"] == "STRUCTURE" for item in pending)


def test_fake_evidence_is_discarded_and_cannot_confirm_pass():
    class Indicator:
        title = "Актуальность"
        configuration = {
            "rule_code": "CD2.R01", "source_document": "source", "source_section": "2.4",
            "source_pages": [14, 15], "capability": "TEXT", "normative_strength": "REQUIRED", "score_weight": 1,
        }

    class Criterion:
        number = "CD2"
        indicators = [Indicator()]

    class Method:
        criteria = [Criterion()]

    output = CandidateAgentOutput(
        criterion_code="CD2", summary="test",
        rule_results=[RuleCheckResult(
            rule_code="CD2.R01", criterion_code="CD2", title="x", status="PASS", finding="x",
            recommendation=None, confidence=1, evidence=[RuleEvidence(quote="выдуманная цитата")],
            source_document="x", source_section="x", source_pages=[1], capability="TEXT",
            normative_strength="REQUIRED", score_weight=1,
        )],
    )
    sanitized = CandidateDissertationAgentFlow(None)._sanitize_output(output, Method(), "Реальный текст документа", "pdf", {"CD2"})
    assert sanitized.rule_results[0].evidence == []
    assert sanitized.rule_results[0].status == "NOT_CHECKED"


def test_docx_extraction_contains_only_available_formatting_metadata(tmp_path):
    source = tmp_path / "candidate.docx"
    extracted = tmp_path / "content.json"
    doc = DocxDocument()
    paragraph = doc.add_paragraph("Введение")
    paragraph.style = "Heading 1"
    paragraph.runs[0].bold = True
    paragraph.runs[0].font.name = "Times New Roman"
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Заголовок"
    doc.save(source)
    payload = TextExtractionService().extract("doc", source, ".docx", extracted)
    assert payload["paragraphs"][0]["formatting"]["runs"][0]["font"] == "Times New Roman"
    assert payload["document_formatting"]["tables"][0]["row_count"] == 2
    assert payload["document_formatting"]["tables"][0]["empty_cells"]
    assert payload["document_formatting"]["pagination_reliable"] is False


@pytest.mark.asyncio
async def test_candidate_agents_run_in_parallel_and_final_receives_compact_package(tmp_path):
    class FakeClient:
        def __init__(self):
            self.active = 0
            self.max_active = 0
            self.final_prompt = ""

        async def ask(self, model, system_prompt, user_prompt, response_model, **kwargs):
            self.active += 1
            self.max_active = max(self.max_active, self.active)
            await asyncio.sleep(0.02)
            self.active -= 1
            if response_model is CandidateFinalOutput:
                self.final_prompt = user_prompt
                output = CandidateFinalOutput(summary="Предварительный анализ завершён.", strengths=[], key_findings=[], recommendations=[])
            else:
                raw = user_prompt.split("Назначенные rules:\n", 1)[1].split("\n\nНедоверенный", 1)[0]
                rules = json.loads(raw)
                output = CandidateAgentOutput(
                    criterion_code=",".join(sorted({item["rule_code"].split(".")[0] for item in rules})),
                    summary="Недостаточно данных.",
                    rule_results=[RuleCheckResult(
                        rule_code=item["rule_code"], criterion_code=item["rule_code"].split(".")[0],
                        title=item["title"], status="NOT_CHECKED", finding="Параметр не проверялся автоматически.",
                        recommendation=None, confidence=1, evidence=[], source_document=item["source_document"],
                        source_section=item["source_section"], source_pages=item["source_pages"],
                        capability=item["capability"], normative_strength=item["normative_strength"],
                        score_weight=item["score_weight"],
                    ) for item in rules],
                )
            return LLMResult(
                output=output, requested_model=model, actual_model=model, aggregator="fake", provider="fake",
                temperature=0, max_completion_tokens=kwargs.get("max_completion_tokens"), usage=LLMUsage(), latency_ms=20,
            )

    extracted = tmp_path / "content.json"
    extracted.write_text(json.dumps({
        "full_text": "Введение. Актуальность темы. Цель и задачи исследования. Заключение.",
        "pages": [{"page_number": 1, "text": "Введение. Актуальность темы. Цель и задачи исследования. Заключение.", "blocks": [{"block_index": 0, "text": "Введение. Актуальность темы. Цель и задачи исследования. Заключение.", "bbox": [0, 0, 1, 1]}]}],
        "page_count": 1,
    }), encoding="utf-8")
    fake = FakeClient()
    async with async_session_factory() as session:
        methodology = await ensure_candidate_dissertation_seed(session)
        document = Document(
            original_name="candidate.pdf", stored_name="candidate.pdf", mime_type="application/pdf", size=1,
            checksum="x", storage_path=str(tmp_path / "candidate.pdf"), extracted_path=str(extracted), extraction_status="completed",
        )
        session.add(document)
        await session.flush()
        analysis = Analysis(
            document_id=document.id, analysis_type="mentor", mode="demo", work_type="CANDIDATE_DISSERTATION",
            methodology_id="CANDIDATE_DISSERTATION", methodology_version="1.0", methodology_name="Кандидатская диссертация",
            methodology_max_score=70, methodology_parameters={}, status="processing", progress=10,
        )
        session.add(analysis)
        await session.flush()
        assessment = Assessment(artifact_type="CANDIDATE_DISSERTATION", artifact_id=document.id, methodology_id=methodology.id, status="processing")
        session.add(assessment)
        await session.commit()
        result = await CandidateDissertationAgentFlow(session, llm_client=fake).execute(assessment, analysis, document)
    assert fake.max_active >= 6
    assert "<document>" not in fake.final_prompt
    assert "not_checked" in fake.final_prompt
    assert result.extra_blocks["nominal_score_max"] == 70
    assert result.extra_blocks["total_score_max"] == 0
    assert all(item.score is None for item in result.criteria)


@pytest.mark.asyncio
async def test_analysis_snapshot_history_and_candidate_report(client):
    async with async_session_factory() as session:
        await ensure_candidate_dissertation_seed(session)
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Введение. Актуальность темы исследования. Цель и задачи исследования.")
    data = pdf.tobytes()
    pdf.close()
    uploaded = await client.post("/api/v1/documents", files={"upload": ("candidate.pdf", io.BytesIO(data), "application/pdf")})
    created = await client.post("/api/v1/analyses", json={"document_id": uploaded.json()["id"], "work_type": "CANDIDATE_DISSERTATION", "mode": "demo"})
    assert created.status_code == 200
    analysis_id = created.json()["analysis_id"]
    async with async_session_factory() as session:
        analysis = await session.get(Analysis, analysis_id)
        assert analysis.methodology_id == "CANDIDATE_DISSERTATION"
        assert analysis.methodology_version == "1.0"
        assert analysis.methodology_max_score == 70
        analysis.status = "completed"
        analysis_result = (
            await session.execute(select(AnalysisResult).where(AnalysisResult.analysis_id == analysis.id))
        ).scalar_one()
        analysis_result.result_json = {
            "overall_score": 8, "criteria": [{"code": "CD1", "title": "Структура", "score": 8, "max_score": 10, "explanation": "Проверено"}],
            "strengths": [], "improvements": [], "remarks": [], "recommendations": [],
            "extra_blocks": {"total_score_max": 10, "candidate_dissertation_report": {
                "evaluated_max_score": 10, "nominal_max_score": 70, "coverage": 0.5,
                    "not_checked": [], "rule_checks": [], "disclaimer": DISCLAIMER,
            }},
        }
        await session.commit()
        document = await session.get(Document, analysis.document_id)
        lines = ReportService()._build_lines(analysis, document, analysis_result.result_json)
        detailed = ReportService()._build_detailed_lines(analysis, document, analysis_result.result_json)
    assert "Предварительный анализ кандидатской диссертации" in lines
    assert "Методология: Кандидатская диссертация" in lines
    assert "Версия: 1.0" in lines
    assert any("предваритель" in line.lower() for line in lines)
    assert "Нормативные основания rule checks:" in detailed
    rendered = fitz.open(stream=ReportService()._render_pdf(lines), filetype="pdf")
    rendered_text = "\n".join(page.get_text() for page in rendered)
    rendered.close()
    assert "Предварительный анализ кандидатской диссертации" in rendered_text
    assert "не являются официальной оценкой" in " ".join(rendered_text.split())
    history = await client.get("/api/v1/analyses/history")
    item = next(value for value in history.json()["items"] if value["analysis_id"] == analysis_id)
    assert item["work_type_display_name"] == "Кандидатская диссертация"
    assert item["methodology_name"] == "Кандидатская диссертация"
    assert item["methodology_version"] == "1.0"
    assert item["total_score_max"] == 10
